# -*- coding: utf-8 -*-
"""生成简约现代风汇报 PPT（16:9，15 页）。"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "export"
OUT = ROOT / "BettaFish-项目汇报.pptx"

INK = RGBColor(0x0F, 0x17, 0x2A)
INK2 = RGBColor(0x47, 0x55, 0x69)
INK3 = RGBColor(0x94, 0xA3, 0xB8)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
SOFT = RGBColor(0xF8, 0xFA, 0xFC)
SOFT2 = RGBColor(0xF1, 0xF5, 0xF9)
ACCENT = RGBColor(0x25, 0x63, 0xEB)
ACCENT_SOFT = RGBColor(0xEF, 0xF6, 0xFF)
BLUE_LT = RGBColor(0xBF, 0xDB, 0xFE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DEEP = RGBColor(0x1D, 0x4E, 0xD8)
FONT = "微软雅黑"

W, H = 13.333, 7.5
M = 0.78


def _ea(run, name=FONT):
    rPr = run.font._rPr
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def txt(slide, x, y, w, h, text, size=14, bold=False, color=INK,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = FONT
        _ea(r)
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.055):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    return shp


def bar(slide, x, y, w=0.62, h=0.045, color=ACCENT):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                               Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False


def page(prs, eyebrow=None, title=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    if eyebrow:
        txt(slide, M, 0.58, 9.0, 0.26, eyebrow, size=10.5, bold=True, color=INK3)
    if title:
        txt(slide, M, 0.9, 11.6, 0.62, title, size=30, bold=True, color=INK)
        bar(slide, M, 1.66)
    return slide, 2.0


def pic(slide, name, y=1.82, max_h=5.12):
    path = IMG / name
    w = max_h * 1.6
    x = (W - w) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                             width=Inches(w), height=Inches(max_h))
    rect(slide, x - 0.012, y - 0.012, w + 0.024, max_h + 0.024,
         fill=None, line=LINE, line_w=0.9, radius=0.02)


def caption(slide, text, y=7.04):
    txt(slide, M, y, W - 2 * M, 0.3, text, size=11, color=INK3)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

# ============ 1 封面 ============
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = INK
txt(s, M, 1.62, 10.5, 0.3, "MULTI-AGENT PUBLIC OPINION ANALYSIS SYSTEM",
    size=11, bold=True, color=RGBColor(0x7D, 0xD3, 0xFC))
txt(s, M, 2.0, 11.0, 1.0, "微舆 BettaFish", size=52, bold=True, color=WHITE)
txt(s, M, 3.05, 11.0, 0.5, "多智能体协同的舆情分析系统", size=23,
    color=RGBColor(0xCB, 0xD5, 0xE1))
bar(s, M, 3.88, w=1.2, h=0.05, color=RGBColor(0x3B, 0x82, 0xF6))
txt(s, M, 4.28, 9.6, 0.9,
    "需求分析与系统设计汇报\n页面原型 · 功能结构 · 角色权限 · 关键技术",
    size=14, color=RGBColor(0x94, 0xA3, 0xB8), spacing=1.6)
txt(s, M, 6.35, 11.2, 0.3,
    "汇报人：（填写姓名）　学号：（填写）　指导教师：（填写）　2026.10",
    size=11, color=RGBColor(0x64, 0x74, 0x8B))
txt(s, W - 2.35, 6.35, 1.6, 0.3, "大三 · 课程设计", size=11,
    color=RGBColor(0x64, 0x74, 0x8B), align=PP_ALIGN.RIGHT)

# ============ 2 背景与痛点 ============
s, y = page(prs, "BACKGROUND", "为什么要做这套系统")
txt(s, M, y - 0.1, 11.8, 0.45,
    "舆情分析的真实工作量，远超「搜一下、写一篇」——痛点集中在三个方面：",
    size=14.5, color=INK2)

pains = [
    ("01", "信息分散，人工检索成本高",
     "同一事件分散在微博、知乎、小红书、新闻网站，逐个平台翻阅效率极低。"),
    ("02", "只有一个大模型，容易「凭空编」",
     "模型知识有截止时间且不能联网，缺少真实数据时仍会给出自信但错误的结论。"),
    ("03", "报告整理耗时，结构难统一",
     "从素材到成稿要经历筛素材、搭框架、配图表，重复劳动多且质量不稳定。"),
]
cw, gap = 3.75, 0.36
for i, (n, t, d) in enumerate(pains):
    x = M + i * (cw + gap)
    rect(s, x, 2.72, cw, 3.0, fill=SOFT, line=LINE, radius=0.05)
    txt(s, x + 0.34, 3.02, 1.0, 0.3, n, size=12, bold=True, color=ACCENT)
    txt(s, x + 0.34, 3.34, cw - 0.68, 0.82, t, size=16, bold=True, color=INK, spacing=1.35)
    txt(s, x + 0.34, 4.18, cw - 0.68, 1.4, d, size=12.5, color=INK2, spacing=1.6)

rect(s, M, 6.02, 11.78, 0.62, fill=ACCENT_SOFT, line=BLUE_LT, radius=0.14)
txt(s, M + 0.32, 6.17, 11.2, 0.35,
    "思路：不指望一个模型搞定，而是组建一支各有分工的 AI 分析团队 —— 让检索、分析、研讨、写作各司其职。",
    size=13, bold=True, color=DEEP)

# ============ 3 系统定位 ============
s, y = page(prs, "POSITIONING", "系统定位：一个舆情分析「工作台」")
txt(s, M, y - 0.12, 11.8, 0.5,
    "输入一个话题，系统自动完成「检索 → 分析 → 研讨 → 成稿」，输出一份可直接交付的结构化报告。",
    size=15, color=INK2)

steps = [("输入", "话题 / 事件名称", "选择一个要分析的对象"),
         ("处理", "多智能体协同", "检索、撰写、交叉点评"),
         ("输出", "结构化报告", "章节 + 图表，可导出")]
for i, (k, t, d) in enumerate(steps):
    x = M + i * 4.06
    rect(s, x, 2.92, 3.62, 1.62, fill=WHITE, line=LINE, radius=0.06)
    txt(s, x + 0.3, 3.12, 1.2, 0.25, k, size=11, bold=True, color=ACCENT)
    txt(s, x + 0.3, 3.48, 3.0, 0.35, t, size=17, bold=True, color=INK)
    txt(s, x + 0.3, 3.95, 3.0, 0.35, d, size=12, color=INK2)
    if i < 2:
        txt(s, x + 3.66, 3.52, 0.4, 0.4, "→", size=17, color=INK3, align=PP_ALIGN.CENTER)

feat = [("不依赖框架", "全部逻辑基于原生 Python 实现，代码可读性强，便于二次开发"),
        ("过程可解释", "每个引擎的检索与思考过程实时可见，结论可回溯"),
        ("数据源可插拔", "搜索服务、模型、数据库均可在配置层替换")]
for i, (t, d) in enumerate(feat):
    x = M + i * 4.06
    bar(s, x, 5.02, w=0.3, h=0.035, color=INK3)
    txt(s, x, 5.22, 3.6, 0.32, t, size=14, bold=True, color=INK)
    txt(s, x, 5.62, 3.55, 0.7, d, size=12, color=INK2, spacing=1.55)

# ============ 4~6 三张图 ============
DIAG = [
    ("07-业务流程图.png", "系统业务流程",
     "用户层 / 调度层 / 智能体层 / 数据层 四条泳道；虚线为核心反思循环"),
    ("08-系统功能结构图.png", "系统功能结构",
     "1 个调度中枢 + 5 大业务模块 + 1 个配置模块，共 18 项细分功能"),
    ("09-角色功能图.png", "角色与功能权限",
     "分析用户 / 系统管理员 / 内容审核员 三类角色，按最小权限原则划分"),
]
for name, title, cap in DIAG:
    s, y = page(prs, "DELIVERABLE", title)
    pic(s, name)
    caption(s, cap)

# ============ 7~9 页面原型 ============
PROTO = [
    ("01-工作台.png", "页面原型 · 分析工作台",
     "统一入口：输入话题、启停引擎、查看运行状态与实时日志"),
    ("02-分析进行中.png", "页面原型 · 分析过程",
     "三引擎日志并排展示，左侧同步显示五阶段进度"),
    ("03-报告详情.png", "页面原型 · 报告阅读",
     "左侧目录导航 + 正文 + 图表，支持导出 HTML / PDF"),
]
for name, title, cap in PROTO:
    s, y = page(prs, "PROTOTYPE", title)
    pic(s, name)
    caption(s, cap)

# ============ 10 技术架构 ============
s, y = page(prs, "ARCHITECTURE", "关键技术 · 分层架构")
layers = [
    ("表现层", "Flask + SocketIO　·　Streamlit × 3　·　HTML/CSS 报告模板", ACCENT_SOFT),
    ("调度层", "子进程管理　·　任务分发　·　WebSocket 实时推送　·　日志采集", SOFT2),
    ("智能体层", "检索节点 → 撰写节点 → 反思节点 → 定稿（三个引擎同构）", WHITE),
    ("能力层", "OpenAI 兼容大模型　·　博查 / Anspire / Tavily　·　本地情感模型", WHITE),
    ("数据层", "PostgreSQL + SQLAlchemy　·　文件产物（Markdown / JSON / HTML）", SOFT2),
]
lh, gap = 0.72, 0.16
for i, (name, tech, fill) in enumerate(layers):
    yy = 2.08 + i * (lh + gap)
    dark = (i == 0)
    rect(s, M, yy, 2.0, lh, fill=INK if dark else SOFT,
         line=None if dark else LINE, radius=0.09)
    txt(s, M, yy + 0.2, 2.0, 0.32, name, size=14, bold=True,
        color=WHITE if dark else INK, align=PP_ALIGN.CENTER)
    rect(s, M + 2.18, yy, 9.6, lh, fill=fill, line=LINE, radius=0.09)
    txt(s, M + 2.52, yy + 0.21, 9.0, 0.32, tech, size=12.5, color=INK2)

caption(s, "上下层单向依赖：上层只依赖下层接口，替换模型或检索服务不影响业务逻辑", 6.94)

# ============ 11 技术选型 ============
s, y = page(prs, "TECH STACK", "关键技术 · 选型说明")
rows = [
    ("大模型接入", "openai SDK", "统一 OpenAI 兼容协议，换厂商只改 Base URL 与模型名"),
    ("Web 服务", "Flask + Flask-SocketIO", "轻量可控，SocketIO 实现分析过程实时推送"),
    ("分析界面", "Streamlit × 3", "每个引擎一个独立应用，按需启动节省内存"),
    ("数据存储", "PostgreSQL + SQLAlchemy", "结构化存储热搜、帖子、评论与报告元信息"),
    ("联网检索", "博查 / Anspire / Tavily", "三种服务可切换，适配不同配额与成本"),
    ("网页采集", "Playwright", "真实浏览器驱动，可处理登录与动态渲染页面"),
    ("本地推理", "PyTorch + Transformers", "多语言情感分类，CPU 即可推理，判断正负面倾向"),
    ("报告渲染", "HTML 模板 + Chart.js", "章节级渲染，图表校验失败自动降级为表格"),
]
ty = 1.98
for h, x, w in [("能力域", M, 2.0), ("技术选型", M + 2.06, 3.5), ("选型理由", M + 5.62, 6.1)]:
    txt(s, x, ty, w, 0.26, h, size=10.5, bold=True, color=INK3)
bar(s, M, ty + 0.32, w=11.78, h=0.008, color=LINE)
for i, (a, b, c) in enumerate(rows):
    ry = ty + 0.46 + i * 0.575
    if i % 2 == 1:
        rect(s, M - 0.1, ry - 0.08, 11.98, 0.53, fill=SOFT, line=None, radius=0.08)
    txt(s, M, ry, 2.0, 0.3, a, size=12, bold=True, color=INK)
    txt(s, M + 2.06, ry, 3.5, 0.3, b, size=12, color=ACCENT)
    txt(s, M + 5.62, ry, 6.1, 0.3, c, size=11.5, color=INK2)

# ============ 12 反思循环 ============
s, y = page(prs, "CORE MECHANISM", "核心机制 · 反思循环（Reflection）")
txt(s, M, y - 0.12, 11.8, 0.4,
    "把「一次搜索就写」改成「写完自己检查、不满意再补搜」，是本系统效果的主要来源。",
    size=14.5, color=INK2)

loop = [("1", "首轮检索", "按话题检索，拿到初始素材"),
        ("2", "撰写初稿", "基于素材生成该段落文本"),
        ("3", "自我检查", "判断还缺什么信息、该搜什么"),
        ("4", "补搜改写", "用新结果改写，循环至达标")]
for i, (n, t, d) in enumerate(loop):
    x = M + i * 3.0
    rect(s, x, 2.78, 2.5, 1.72, fill=WHITE, line=LINE, radius=0.07)
    txt(s, x + 0.24, 2.94, 0.4, 0.28, n, size=11, bold=True, color=ACCENT)
    txt(s, x + 0.24, 3.26, 2.0, 0.32, t, size=15.5, bold=True, color=INK)
    txt(s, x + 0.24, 3.66, 2.05, 0.7, d, size=11.5, color=INK2, spacing=1.5)
    if i < 3:
        txt(s, x + 2.56, 3.3, 0.4, 0.4, "→", size=16, color=INK3, align=PP_ALIGN.CENTER)

rect(s, M, 4.72, 5.9, 1.6, fill=INK, line=None, radius=0.07)
txt(s, M + 0.34, 4.94, 5.2, 1.2,
    "日志实况：\n反思 1/2 …\n  反思查询: 美国队 失误 得分 细节\n  找到 10 个反思结果",
    size=11, color=RGBColor(0xCA, 0xD5, 0xE1), spacing=1.7)

rect(s, M + 6.16, 4.72, 5.62, 1.6, fill=ACCENT_SOFT, line=BLUE_LT, radius=0.07)
txt(s, M + 6.5, 4.94, 5.0, 0.34, "可调参数（配置层即可修改）",
    size=13, bold=True, color=DEEP)
txt(s, M + 6.5, 5.36, 5.0, 0.85,
    "反思轮数：默认 2，越大越深入、耗时与成本越高\n段落数：默认 5，影响报告篇幅",
    size=11.5, color=RGBColor(0x1E, 0x3A, 0x8A), spacing=1.8)

# ============ 13 多智能体协同 ============
s, y = page(prs, "COLLABORATION", "多智能体协同 · 论坛式交叉印证")
txt(s, M, y - 0.12, 11.8, 0.4,
    "三个引擎不是各写各的：它们把阶段性结论放进同一个「论坛」，互相看见对方的判断。",
    size=14.5, color=INK2)

agents = [("InsightAgent", "库内数据视角", "看采集到的帖子与评论"),
          ("MediaAgent", "联网媒介视角", "看新闻报道与图片"),
          ("QueryAgent", "深度检索视角", "看背景与延伸报道")]
for i, (n, role, d) in enumerate(agents):
    x = M + i * 2.62
    rect(s, x, 2.55, 2.3, 1.3, fill=WHITE, line=LINE, radius=0.07)
    txt(s, x + 0.22, 2.7, 1.9, 0.26, n, size=12.5, bold=True, color=INK)
    txt(s, x + 0.22, 3.0, 1.9, 0.24, role, size=11, color=ACCENT)
    txt(s, x + 0.22, 3.3, 1.92, 0.5, d, size=10.5, color=INK2, spacing=1.45)

for i in range(3):
    cx = M + i * 2.62 + 1.15
    s.shapes.add_connector(1, Inches(cx), Inches(3.9), Inches(cx), Inches(4.16)) \
        .line.color.rgb = BLUE_LT

rect(s, 3.4, 4.2, 6.5, 1.34, fill=INK, line=None, radius=0.07)
txt(s, 3.4, 4.38, 6.5, 0.3, "ForumEngine　共享观点池",
    size=13.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
txt(s, 4.0, 4.78, 5.3, 0.6,
    "持续监听三引擎输出 → 生成 HOST 点评 → 写回共享日志 → 被各引擎读回",
    size=11, color=RGBColor(0x94, 0xA3, 0xB8), spacing=1.7, align=PP_ALIGN.CENTER)

txt(s, M, 5.86, 5.6, 0.32, "价值：减少「三个瞎子摸象」", size=13, bold=True, color=INK)
txt(s, M, 6.24, 5.6, 0.9,
    "结论冲突在成稿前就被暴露，报告一致性更高；\n同时保留分歧表述，避免强行统一口径。",
    size=11.5, color=INK2, spacing=1.6)
txt(s, M + 6.16, 5.86, 5.6, 0.32, "工程上怎么做到的", size=13, bold=True, color=INK)
txt(s, M + 6.16, 6.24, 5.6, 0.9,
    "日志监听 + 共享文件：无需引入消息队列，\n靠文件变化感知即可实现跨进程协同。",
    size=11.5, color=INK2, spacing=1.6)

# ============ 14 实测成效 ============
s, y = page(prs, "RESULT", "实测成效（本机真实运行数据）")
kpis = [("213", "条", "单次采集入库热搜", "9 个新闻源成功"),
        ("20", "个", "提炼关键词", "供深度采集使用"),
        ("220", "秒", "单份报告生成耗时", "6 章节 / 5 张图表"),
        ("9", "张", "原型与结构图产出", "含 6 个页面原型")]
for i, (v, u, l, d) in enumerate(kpis):
    x = M + i * 3.02
    rect(s, x, 2.36, 2.82, 1.98, fill=SOFT, line=LINE, radius=0.06)
    txt(s, x + 0.3, 2.6, 1.6, 0.72, v, size=36, bold=True, color=ACCENT)
    txt(s, x + 1.5, 3.14, 0.6, 0.3, u, size=13, color=INK3)
    txt(s, x + 0.3, 3.44, 2.3, 0.3, l, size=13, bold=True, color=INK)
    txt(s, x + 0.3, 3.82, 2.3, 0.32, d, size=11, color=INK2)

txt(s, M, 4.72, 5.8, 0.32, "验证结论", size=13.5, bold=True, color=INK)
txt(s, M, 5.1, 5.8, 1.5,
    "① 完整链路可跑通：采集 → 分析 → 成稿全程无人工干预\n"
    "② 数据源决定质量：无真实数据时模型会编造结论\n"
    "③ 单份报告成本可控：约 12 次检索 + 数十次模型调用",
    size=11.5, color=INK2, spacing=1.9)
txt(s, M + 6.3, 4.72, 5.5, 0.32, "已知不足", size=13.5, bold=True, color=INK)
txt(s, M + 6.3, 5.1, 5.5, 1.5,
    "① 依赖第三方检索配额，超额后分析降级\n"
    "② 深度采集需人工登录平台，暂不能无人值守\n"
    "③ PDF 导出依赖系统级图形库，跨平台需额外配置",
    size=11.5, color=INK2, spacing=1.9)

# ============ 15 总结 ============
s, y = page(prs, "SUMMARY", "总结与展望")
txt(s, M, y - 0.12, 11.8, 0.4,
    "本阶段完成「需求分析 + 系统设计 + 原型验证」三部分工作。", size=14.5, color=INK2)

done = [("需求与原型", "梳理 6 个核心页面，输出可交互原型与高清图"),
        ("结构与权限", "输出功能结构图与三类角色的功能权限划分"),
        ("技术验证", "完成技术选型，并在本机跑通完整业务链路")]
for i, (t, d) in enumerate(done):
    x = M + i * 4.06
    rect(s, x, 2.66, 3.64, 1.34, fill=ACCENT_SOFT, line=BLUE_LT, radius=0.07)
    txt(s, x + 0.28, 2.84, 3.1, 0.32, t, size=14, bold=True, color=DEEP)
    txt(s, x + 0.28, 3.22, 3.12, 0.66, d, size=11.5,
        color=RGBColor(0x1E, 0x3A, 0x8A), spacing=1.5)

txt(s, M, 4.42, 5.8, 0.32, "下一步", size=13.5, bold=True, color=INK)
txt(s, M, 4.8, 5.8, 1.5,
    "· 增加多级审批与报告版本管理\n· 引入检索结果缓存，降低调用成本\n· 支持定时任务与订阅式舆情监控",
    size=11.5, color=INK2, spacing=1.9)
txt(s, M + 6.3, 4.42, 5.5, 0.32, "个人收获", size=13.5, bold=True, color=INK)
txt(s, M + 6.3, 4.8, 5.5, 1.5,
    "· 掌握多智能体系统的基本设计范式\n· 理解「数据质量决定模型输出」的工程规律\n· 具备从需求到原型的完整表达训练",
    size=11.5, color=INK2, spacing=1.9)

rect(s, M, 6.54, 11.78, 0.5, fill=INK, line=None, radius=0.16)
txt(s, M, 6.64, 11.78, 0.32, "感谢观看　·　请老师批评指正", size=13, bold=True,
    color=WHITE, align=PP_ALIGN.CENTER)

prs.save(OUT)
print("已生成:", OUT)
print("页数:", len(prs.slides._sldIdLst))
