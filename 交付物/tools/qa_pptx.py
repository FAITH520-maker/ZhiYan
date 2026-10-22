# -*- coding: utf-8 -*-
"""PPT 自检：形状越界、图片数量、文字溢出估算。"""
import math
import unicodedata
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
PPTX = ROOT / "BettaFish-项目汇报.pptx"

EMU_IN = 914400


def text_width_pt(s, size):
    """粗估一行文字宽度（pt）：中日韩字符≈1em，其他≈0.55em。"""
    w = 0.0
    for ch in s:
        if unicodedata.east_asian_width(ch) in ("W", "F"):
            w += size
        else:
            w += size * 0.55
    return w


def check():
    prs = Presentation(str(PPTX))
    sw, sh = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    issues = 0

    for idx, slide in enumerate(prs.slides, 1):
        pics = 0
        warn = []
        for shp in slide.shapes:
            if shp.shape_type == 13:  # PICTURE
                pics += 1
            if shp.left is None:
                continue
            r = (shp.left + shp.width) / EMU_IN
            b = (shp.top + shp.height) / EMU_IN
            if r > sw + 0.02 or b > sh + 0.02 or shp.left < -0.02 or shp.top < -0.02:
                warn.append(f"越界 {shp.shape_type} 至 ({r:.2f},{b:.2f})")

            if not shp.has_text_frame:
                continue
            box_w = shp.width / EMU_IN - 0.02
            box_h = shp.height / EMU_IN
            total = 0.0
            for p in shp.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs)
                if not text:
                    continue
                size = None
                for r in p.runs:
                    if r.font.size:
                        size = r.font.size.pt
                        break
                size = size or 14.0
                lines = max(1, math.ceil(text_width_pt(text, size) / (box_w * 72)))
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
                total += lines * size * ls * 1.25 / 72
            if total > box_h * 1.18 and total - box_h > 0.12:
                first = "".join(r.text for r in shp.text_frame.paragraphs[0].runs)[:18]
                warn.append(f"文字可能溢出 需{total:.2f}\" 框高{box_h:.2f}\" 「{first}」")

        tag = "OK" if not warn else "!!"
        if warn:
            issues += 1
        print(f"[{tag}] 第 {idx:2d} 页  形状 {len(slide.shapes):2d}  图片 {pics}  "
              + ("；".join(warn[:3]) if warn else ""))

    print(f"\n共 {len(prs.slides._sldIdLst)} 页，画布 {sw:.2f}x{sh:.2f} in，问题页 {issues}")
    print("文件大小:", round(PPTX.stat().st_size / 1024 / 1024, 2), "MB")


if __name__ == "__main__":
    check()
